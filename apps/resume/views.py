# from django.shortcuts import get_object_or_404
# from rest_framework import viewsets
# from django.http import JsonResponse
# import pdfplumber
# import re
# import os
# from docx import Document
# from apps.candidate.models import Candidates
# from apps.resume.models import ResumeFile
# from apps.resume.serializers import ResumeFileSerializer
# 
# 
# class ResumeFileViewSet(viewsets.ModelViewSet):
#     queryset = ResumeFile.objects.all()
#     serializer_class = ResumeFileSerializer
# 
# 
# def extract_details(text):
#     email_regex = r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
#     phone_regex = r"\+?\d{1,3}[-.\s]?\(?\d{2,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}"
# 
#     emails = re.findall(email_regex, text)
#     phones = re.findall(phone_regex, text)
# 
#     email = emails[0] if emails else None
#     phone = phones[0] if phones else None
# 
#     name = None
#     lines = text.split("\n")
#     for line in lines[:5]:  # Checking the first 5 lines for a name
#         if len(line.split()) >= 2 and all(word[0].isupper() for word in line.split()):
#             name = line.strip()
#             break
# 
#     return {"name": name, "emails": emails, "phones": phones}
# 
# 
# def parse_file(request, file_id):
#     try:
#         uploaded_file = get_object_or_404(ResumeFile, id=file_id)
#         file_path = uploaded_file.file.path
#         file_extension = os.path.splitext(file_path)[1].lower()
# 
#         text = ""
# 
#         if file_extension == ".pdf":
#             with pdfplumber.open(file_path) as pdf:
#                 for page in pdf.pages:
#                     extracted_text = page.extract_text()
#                     if extracted_text:
#                         text += extracted_text + "\n"
# 
#         elif file_extension == ".docx":
#             doc = Document(file_path)
#             for para in doc.paragraphs:
#                 text += para.text + "\n"
# 
#         else:
#             return JsonResponse({"error": "Unsupported file type"}, status=400)
# 
#         extracted_data = extract_details(text)
# 
#         # Stores parsed details in Candidate Model
#         candidates, created = Candidates.objects.update_or_create(
#             resume=uploaded_file,
#             defaults={
#                 "name": extracted_data["name"],
#                 "email": extracted_data["email"],
#                 "phone": extracted_data["phone"],
#             },
#         )
# 
#         # return JsonResponse({"text": text, "extracted_data": extracted_data})
#         return JsonResponse({
#             "text": text,
#             "extracted_data": extracted_data,
#             "candidate_id": Candidates.id,
#             "updated": not created  # True if existing record was updated
#         })
# 
#     except Exception as e:
#         return JsonResponse({"error": str(e)}, status=500)
# 



from django.shortcuts import get_object_or_404
from rest_framework import viewsets
from django.http import JsonResponse
import pdfplumber
import re
import os
from docx import Document
from apps.resume.models import ResumeFile
from apps.resume.serializers import ResumeFileSerializer
from apps.candidate.models import Candidate  # Import Candidate model


class ResumeFileViewSet(viewsets.ModelViewSet):
    queryset = ResumeFile.objects.all()
    serializer_class = ResumeFileSerializer


def extract_details(text):
    email_regex = r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}"
    phone_regex = r"\+?\d{1,3}[-.\s]?\(?\d{2,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}"


    emails = re.findall(email_regex, text)
    phones = re.findall(phone_regex, text)

    # Extract first valid email & phone
    email = emails[0] if emails else None
    phone = phones[0] if phones else None

    name = None
    lines = text.split("\n")
    for line in lines[:5]:  # Checking the first 5 lines for a name
        if len(line.split()) >= 2 and all(word[0].isupper() for word in line.split()):
            name = line.strip()
            break

    return {"name": name, "email": email, "phone": phone}


def parse_file(request, file_id):
    try:
        uploaded_file = get_object_or_404(ResumeFile, id=file_id)
        file_path = uploaded_file.file.path
        file_extension = os.path.splitext(file_path)[1].lower()

        text = ""

        if file_extension == ".pdf":
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    extracted_text = page.extract_text()
                    if extracted_text:
                        text += extracted_text + "\n"

        elif file_extension == ".docx":
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"

        else:
            return JsonResponse({"error": "Unsupported file type"}, status=400)

        extracted_data = extract_details(text)

        # Store parsed details in Candidate model
        candidate, created = Candidate.objects.update_or_create(
            resume=uploaded_file,
            defaults={
                "name": extracted_data["name"],
                "email": extracted_data["email"],
                "phone": extracted_data["phone"],
            },
        )

        return JsonResponse({
            "text": text,
            "extracted_data": extracted_data,
            "candidate_id": candidate.id,
            "updated": not created  # True if existing record was updated
        })

    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)
